import hashlib
import hmac
import logging
import os
import secrets
from contextlib import contextmanager, nullcontext
from datetime import datetime
from functools import wraps
from io import BytesIO
from pathlib import Path

import mysql.connector
from flask import Flask, g, jsonify, request, send_file, send_from_directory
from flask_cors import CORS
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash

load_dotenv()

from db import connection


class ApiError(Exception):
    """Erro conhecido da API, sempre serializado como JSON."""

    def __init__(self, message, status=400):
        self.message = message
        self.status = status
        super().__init__(message)


app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("APP_SECRET_KEY") or secrets.token_urlsafe(32)

# O front-end pode ser servido pelo `npx serve` (3000), Live Server (5500) ou
# pelo próprio Flask. A lista pode ser restringida com CORS_ORIGINS no .env.
cors_setting = os.getenv(
    "CORS_ORIGINS",
    "http://localhost:3000,http://127.0.0.1:3000,"
    "http://localhost:5500,http://127.0.0.1:5500,"
    "http://localhost:5173,http://127.0.0.1:5173",
).strip()
cors_origins = "*" if cors_setting == "*" else [
    origin.strip() for origin in cors_setting.split(",") if origin.strip()
]
CORS(
    app,
    resources={r"/api/*": {"origins": cors_origins}},
    methods=["GET", "POST", "PATCH", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
    expose_headers=["Content-Disposition"],
)


@app.get("/")
def frontend_index():
    """Serve a interface no mesmo endereço da API para evitar CORS em produção."""
    return send_from_directory(app.root_path, "index.html")


@app.get("/css/<path:filename>")
def frontend_css(filename):
    return send_from_directory(os.path.join(app.root_path, "css"), filename)


@app.get("/js/<path:filename>")
def frontend_js(filename):
    return send_from_directory(os.path.join(app.root_path, "js"), filename)


@app.get("/assets/<path:filename>")
def frontend_assets(filename):
    return send_from_directory(os.path.join(app.root_path, "assets"), filename)

token_serializer = URLSafeTimedSerializer(app.config["SECRET_KEY"], salt="fila-auditoria-auth")
TOKEN_MAX_AGE_SECONDS = int(os.getenv("TOKEN_MAX_AGE_SECONDS", "28800"))
QUEUE_LOCK_NAME = "fila_auditoria_mutation"
QUEUE_LOCK_TIMEOUT_SECONDS = 5
LOGO_PATH = Path(app.root_path) / "assets" / "logo-unicomgroup.png"

TIM_BLUE = "001A9C"
TIM_RED = "E2001A"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
METRIC_PERIODS = {
    "dia": ("Hoje", "DATE(ac.data_conclusao) = CURRENT_DATE()"),
    "semana": (
        "Últimos 7 dias",
        "ac.data_conclusao >= DATE_SUB(CURRENT_DATE(), INTERVAL 7 DAY) "
        "AND ac.data_conclusao < DATE_ADD(CURRENT_DATE(), INTERVAL 1 DAY)",
    ),
    "mes": (
        "Mês atual",
        "YEAR(ac.data_conclusao) = YEAR(CURRENT_DATE()) "
        "AND MONTH(ac.data_conclusao) = MONTH(CURRENT_DATE())",
    ),
}


def error(message, status=400):
    return jsonify({"erro": message}), status


@app.errorhandler(ApiError)
def handle_api_error(exc):
    return error(exc.message, exc.status)


@app.errorhandler(HTTPException)
def handle_http_error(exc):
    return error(exc.description, exc.code or 500)


@app.errorhandler(mysql.connector.Error)
def handle_database_error(exc):
    app.logger.exception("Erro de banco de dados")
    return error("Não foi possível acessar o banco de dados.", 503)


@app.errorhandler(Exception)
def handle_unexpected_error(exc):
    app.logger.exception("Erro inesperado")
    return error("Ocorreu um erro interno.", 500)


def request_data():
    """Exige um objeto JSON; evita que uma resposta HTML/None quebre o fetch."""
    if not request.is_json:
        raise ApiError("Envie o corpo da requisição como JSON (Content-Type: application/json).", 415)
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ApiError("O JSON da requisição é inválido.", 400)
    return data


def required_text(data, field, label, max_length=100):
    value = data.get(field)
    if not isinstance(value, str) or not value.strip():
        raise ApiError(f"{label} é obrigatório.")
    value = " ".join(value.strip().split())
    if len(value) > max_length:
        raise ApiError(f"{label} deve ter no máximo {max_length} caracteres.")
    return value


def optional_text(data, field, label, max_length=100):
    if field not in data:
        return None
    return required_text(data, field, label, max_length)


def password_from(data, field="senha", required=True):
    if field not in data and not required:
        return None
    password = data.get(field)
    if not isinstance(password, str) or not password:
        raise ApiError("Senha é obrigatória.")
    if len(password) > 128:
        raise ApiError("Senha deve ter no máximo 128 caracteres.")
    return password


def user_type_from(data, required=True):
    if "tipo_usuario" not in data and not required:
        return None
    user_type = data.get("tipo_usuario")
    if not isinstance(user_type, str):
        raise ApiError("Tipo de usuário é obrigatório.")
    # AUDITOR é aceito somente para facilitar a migração do cliente anterior.
    aliases = {"AUDITOR": "FUNCIONARIO", "FUNCIONARIO": "FUNCIONARIO", "GESTOR": "GESTOR"}
    normalized = aliases.get(user_type.strip().upper())
    if not normalized:
        raise ApiError("Tipo de usuário deve ser GESTOR ou FUNCIONARIO.")
    return normalized


def public_user(user):
    return {
        "id": user["id"],
        "nome": user["nome"],
        "tipo_usuario": user["tipo_usuario"],
        "ativo": bool(user.get("ativo", True)),
    }


def verify_password(stored_password, supplied_password):
    """Aceita temporariamente SHA-256 do schema antigo e o migra após login."""
    if stored_password.startswith(("scrypt:", "pbkdf2:")):
        return check_password_hash(stored_password, supplied_password), False
    legacy_hash = hashlib.sha256(supplied_password.encode("utf-8")).hexdigest()
    return hmac.compare_digest(stored_password, legacy_hash), True


def auth_token(user):
    return token_serializer.dumps({"usuario_id": user["id"]})


def token_from_authorization():
    authorization = request.headers.get("Authorization", "")
    if not authorization.startswith("Bearer "):
        raise ApiError("Autenticação obrigatória.", 401)
    return authorization.removeprefix("Bearer ").strip()


def current_user_from_token(token):
    """Valida o token e devolve somente um usuário atualmente ativo."""
    if not isinstance(token, str) or not token:
        raise ApiError("Sessão inválida ou expirada. Entre novamente.", 401)
    try:
        payload = token_serializer.loads(token, max_age=TOKEN_MAX_AGE_SECONDS)
        user_id = int(payload["usuario_id"])
    except (BadSignature, SignatureExpired, KeyError, TypeError, ValueError):
        raise ApiError("Sessão inválida ou expirada. Entre novamente.", 401) from None

    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT id, nome, tipo_usuario, ativo FROM usuarios WHERE id = %s AND ativo = TRUE",
            (user_id,),
        )
        user = cursor.fetchone()
    if not user:
        raise ApiError("Usuário inativo ou não encontrado.", 401)
    return user


def logout_token_from_request():
    """Aceita Bearer normal ou o token enviado pelo navigator.sendBeacon."""
    authorization = request.headers.get("Authorization", "")
    if authorization.startswith("Bearer "):
        return authorization.removeprefix("Bearer ").strip()

    data = request.get_json(silent=True)
    token = data.get("token") if isinstance(data, dict) else None
    if not isinstance(token, str) or not token:
        raise ApiError("Autenticação obrigatória.", 401)
    return token


def authenticated(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        g.current_user = current_user_from_token(token_from_authorization())
        return view(*args, **kwargs)

    return wrapper


def manager_required(view):
    @authenticated
    @wraps(view)
    def wrapper(*args, **kwargs):
        if g.current_user["tipo_usuario"] != "GESTOR":
            raise ApiError("Apenas gestores podem executar esta ação.", 403)
        return view(*args, **kwargs)

    return wrapper


@contextmanager
def queue_mutation_lock(conn):
    """Serializa mudanças na fila para manter posições únicas sob concorrência."""
    lock_cursor = conn.cursor()
    acquired = False
    try:
        lock_cursor.execute("SELECT GET_LOCK(%s, %s)", (QUEUE_LOCK_NAME, QUEUE_LOCK_TIMEOUT_SECONDS))
        acquired = lock_cursor.fetchone()[0] == 1
        if not acquired:
            raise ApiError("A fila está sendo atualizada. Tente novamente em instantes.", 409)
        yield
    finally:
        if acquired:
            lock_cursor.execute("SELECT RELEASE_LOCK(%s)", (QUEUE_LOCK_NAME,))
            lock_cursor.fetchone()
        lock_cursor.close()


def queue_rows(cursor):
    cursor.execute(
        """SELECT f.id, f.usuario_id, u.nome, f.posicao, f.status
           FROM fila_auditoria f
           JOIN usuarios u ON u.id = f.usuario_id
           WHERE u.ativo = TRUE
           ORDER BY f.posicao"""
    )
    return [
        {
            "id": row[0],
            "usuario_id": row[1],
            "nome": row[2],
            "posicao": row[3],
            "status": row[4],
        }
        for row in cursor.fetchall()
    ]


def find_user_for_queue(cursor, data):
    user_id = data.get("usuario_id")
    if user_id:
        cursor.execute(
            "SELECT id, nome, tipo_usuario FROM usuarios WHERE id = %s AND ativo = TRUE", (user_id,)
        )
    else:
        name = required_text(data, "nome", "Nome")
        cursor.execute(
            "SELECT id, nome, tipo_usuario FROM usuarios WHERE LOWER(nome) = LOWER(%s) AND ativo = TRUE",
            (name,),
        )
    return cursor.fetchone()


def set_queue_order(cursor, item_ids):
    """Grava a ordem recebida sem violar a chave única de posição."""
    if not item_ids:
        return
    # A inversão temporária evita colisão na chave única de posicao.
    cursor.execute("UPDATE fila_auditoria SET posicao = -posicao")
    for position, item_id in enumerate(item_ids, 1):
        cursor.execute("UPDATE fila_auditoria SET posicao = %s WHERE id = %s", (position, item_id))
    cursor.execute("UPDATE fila_auditoria SET status = 'AGUARDANDO'")
    cursor.execute("UPDATE fila_auditoria SET status = 'EM_ANDAMENTO' WHERE id = %s", (item_ids[0],))


def normalize_queue(cursor):
    cursor.execute("SELECT id FROM fila_auditoria ORDER BY posicao, id")
    item_ids = [row[0] for row in cursor.fetchall()]
    set_queue_order(cursor, item_ids)


def enqueue_employee(cursor, user_id):
    """Inclui o funcionário ativo no fim da fila; a operação é idempotente."""
    cursor.execute(
        "SELECT tipo_usuario, ativo FROM usuarios WHERE id = %s FOR UPDATE", (user_id,)
    )
    user = cursor.fetchone()
    if not user:
        return False
    user_type = user["tipo_usuario"] if isinstance(user, dict) else user[0]
    active = user["ativo"] if isinstance(user, dict) else user[1]
    if user_type != "FUNCIONARIO" or not active:
        return False

    cursor.execute("SELECT id FROM fila_auditoria WHERE usuario_id = %s FOR UPDATE", (user_id,))
    if cursor.fetchone():
        return False

    cursor.execute("SELECT posicao FROM fila_auditoria ORDER BY posicao DESC LIMIT 1 FOR UPDATE")
    last_item = cursor.fetchone()
    last_position = (last_item["posicao"] if isinstance(last_item, dict) else last_item[0]) if last_item else 0
    position = last_position + 1
    cursor.execute(
        "INSERT INTO fila_auditoria (usuario_id, posicao, status) VALUES (%s, %s, %s)",
        (user_id, position, "EM_ANDAMENTO" if position == 1 else "AGUARDANDO"),
    )
    return True


def remove_employee_from_queue(cursor, user_id):
    """Remove um funcionário da fila e recalcula a vez atual, se necessário."""
    cursor.execute("DELETE FROM fila_auditoria WHERE usuario_id = %s", (user_id,))
    removed = cursor.rowcount > 0
    if removed:
        normalize_queue(cursor)
    return removed


def metric_period(value):
    period = (value or "dia").strip().lower()
    if period not in METRIC_PERIODS:
        raise ApiError("Período inválido. Use dia, semana ou mes.")
    return period, *METRIC_PERIODS[period]


def metrics_payload(cursor, period):
    """Total e ranking consolidados no período solicitado."""
    period, label, condition = metric_period(period)
    cursor.execute(f"SELECT COUNT(*) AS total FROM auditorias_concluidas ac WHERE {condition}")
    total = cursor.fetchone()["total"]
    cursor.execute(
        f"""SELECT u.id AS usuario_id, u.nome, COUNT(ac.id) AS total
            FROM usuarios u
            LEFT JOIN auditorias_concluidas ac
              ON ac.usuario_id = u.id AND {condition}
            WHERE u.tipo_usuario = 'FUNCIONARIO'
            GROUP BY u.id, u.nome
            ORDER BY total DESC, u.nome"""
    )
    return {"periodo": period, "titulo": label, "total_auditorias": total, "ranking": cursor.fetchall()}


def audit_ranking_payload(cursor):
    """Ranking geral de auditorias dos funcionários atualmente ativos."""
    cursor.execute(
        """SELECT u.id AS usuario_id, u.nome, COUNT(ac.id) AS total
           FROM usuarios u
           LEFT JOIN auditorias_concluidas ac ON ac.usuario_id = u.id
           WHERE u.tipo_usuario = 'FUNCIONARIO' AND u.ativo = TRUE
           GROUP BY u.id, u.nome
           HAVING COUNT(ac.id) > 0
           ORDER BY total DESC, u.nome"""
    )
    return cursor.fetchall()


def history_rows(cursor, period):
    period, label, condition = metric_period(period)
    cursor.execute(
        f"""SELECT u.nome AS auditor, ac.data_conclusao
            FROM auditorias_concluidas ac
            JOIN usuarios u ON u.id = ac.usuario_id
            WHERE {condition}
            ORDER BY ac.data_conclusao DESC, u.nome"""
    )
    return label, cursor.fetchall()


def brand_report_header(sheet, title, subtitle, end_column):
    if not LOGO_PATH.is_file():
        raise ApiError("A logomarca do relatório não foi encontrada em assets/logo-unicomgroup.png.", 500)
    for row in range(1, 4):
        sheet.row_dimensions[row].height = 24
    logo = OpenpyxlImage(str(LOGO_PATH))
    logo.width = 260
    logo.height = 49
    sheet.add_image(logo, "A1")
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=end_column)
    sheet.merge_cells(start_row=6, start_column=1, end_row=6, end_column=end_column)
    title_cell = sheet.cell(5, 1, title)
    title_cell.font = Font(name="Arial", size=14, bold=True, color=TIM_BLUE)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    subtitle_cell = sheet.cell(6, 1, subtitle)
    subtitle_cell.font = Font(name="Arial", size=11, italic=True, color="4B5563")
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[5].height = 25
    sheet.row_dimensions[6].height = 20
    sheet.sheet_view.showGridLines = False


def write_report_table(sheet, headers, rows):
    header_row = 8
    blue_fill = PatternFill("solid", fgColor=TIM_BLUE)
    red_side = Side(style="thin", color=TIM_RED)
    border = Border(left=red_side, right=red_side, top=red_side, bottom=red_side)
    centered = Alignment(horizontal="center", vertical="center")
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = blue_fill
        cell.font = Font(name="Arial", size=12, bold=True, color=WHITE)
        cell.alignment = centered
        cell.border = border
    for row_index, values in enumerate(rows, header_row + 1):
        is_zebra_row = (row_index - header_row) % 2 == 0
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="Arial", size=11)
            cell.alignment = centered
            cell.border = border
            if is_zebra_row:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
            if isinstance(value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
    last_row = max(header_row, header_row + len(rows))
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    sheet.freeze_panes = "A9"
    for column in range(1, len(headers) + 1):
        widest = max(len(str(sheet.cell(row, column).value or "")) for row in range(5, last_row + 1))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(widest + 3, 16), 45)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def build_excel_report(period):
    workbook = Workbook()
    history_sheet = workbook.active
    history_sheet.title = "Relatório de Auditorias"
    summary_sheet = workbook.create_sheet("Resumo de Produtividade")
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        metrics = metrics_payload(cursor, period)
        label, history = history_rows(cursor, period)

    report_date = datetime.now().strftime("%d/%m/%Y %H:%M")
    brand_report_header(
        history_sheet,
        "Relatório de Auditorias",
        f"Período: {label} | Gerado em: {report_date}",
        3,
    )
    history_data = [(row["auditor"], row["data_conclusao"], label) for row in history]
    if not history_data:
        history_data = [("Nenhuma auditoria concluída no período", "", label)]
    write_report_table(history_sheet, ["Nome do Auditor", "Data/Hora da Conclusão", "Período"], history_data)

    brand_report_header(
        summary_sheet,
        "Resumo de Produtividade",
        f"Período: {metrics['titulo']} | Total geral: {metrics['total_auditorias']} auditoria(s)",
        4,
    )
    ranking_data = [
        (position, item["nome"], item["total"], metrics["titulo"])
        for position, item in enumerate(metrics["ranking"], 1)
    ]
    if not ranking_data:
        ranking_data = [("-", "Nenhum funcionário cadastrado", 0, metrics["titulo"])]
    write_report_table(
        summary_sheet,
        ["Posição", "Auditor", "Auditorias Concluídas", "Período"],
        ranking_data,
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
    for item_id in ids:
        cursor.execute("UPDATE fila_auditoria SET posicao = -posicao WHERE id = %s", (item_id,))
    for position, item_id in enumerate(ids, 1):
        cursor.execute("UPDATE fila_auditoria SET posicao = %s WHERE id = %s", (position, item_id))
    if ids:
        cursor.execute(
            "UPDATE fila_auditoria SET status = CASE WHEN posicao = 1 "
            "THEN 'EM_ANDAMENTO' ELSE 'AGUARDANDO' END"
        )


@app.get("/api/health")
def health():
    """Teste rápido para separar erro de CORS/rota de erro de banco."""
    try:
        with connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT 1")
            cursor.fetchone()
        return jsonify({"status": "ok"})
    except mysql.connector.Error:
        return jsonify({"status": "indisponivel", "erro": "Banco de dados inacessível."}), 503


@app.route("/api/login", methods=["POST"])
def login():
    data = request_data()
    # `usuario` é mantido como alias para clientes que ainda não foram atualizados.
    name = data.get("nome", data.get("usuario"))
    name = required_text({"nome": name}, "nome", "Nome de usuário")
    password = password_from(data)

    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT id, nome, senha, tipo_usuario, ativo FROM usuarios "
            "WHERE LOWER(nome) = LOWER(%s)",
            (name,),
        )
        user = cursor.fetchone()
        valid_password = False
        legacy_password = False
        if user and user["ativo"]:
            valid_password, legacy_password = verify_password(user["senha"], password)
        if not valid_password:
            raise ApiError("Nome de usuário ou senha inválidos.", 401)
        if legacy_password:
            cursor.execute(
                "UPDATE usuarios SET senha = %s WHERE id = %s",
                (generate_password_hash(password), user["id"]),
            )
        joined_queue = False
        if user["tipo_usuario"] == "FUNCIONARIO":
            # A chave única por usuário e a trava da fila tornam o login
            # idempotente, inclusive quando há duas tentativas simultâneas.
            with queue_mutation_lock(conn):
                joined_queue = enqueue_employee(cursor, user["id"])
        conn.commit()

    return jsonify(
        {
            "usuario": public_user(user),
            "token": auth_token(user),
            "entrou_na_fila": joined_queue,
        }
    )


@app.post("/api/logout")
def logout():
    """Encerra a participação do funcionário inclusive em chamadas sendBeacon."""
    user = current_user_from_token(logout_token_from_request())
    removed_from_queue = False
    if user["tipo_usuario"] == "FUNCIONARIO":
        with connection() as conn:
            cursor = conn.cursor()
            with queue_mutation_lock(conn):
                removed_from_queue = remove_employee_from_queue(cursor, user["id"])
            conn.commit()
    return jsonify(
        {
            "mensagem": "Sessão encerrada com sucesso.",
            "removido_da_fila": removed_from_queue,
        }
    )


@app.post("/api/fila/entrar")
@authenticated
def join_own_queue():
    """Reinsere apenas o próprio funcionário ao restaurar uma aba recarregada."""
    if g.current_user["tipo_usuario"] != "FUNCIONARIO":
        raise ApiError("Apenas funcionários participam da fila.", 409)
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            joined_queue = enqueue_employee(cursor, g.current_user["id"])
        conn.commit()
        return jsonify({"entrou_na_fila": joined_queue, "fila": queue_rows(cursor)})


@app.patch("/api/minha-senha")
@authenticated
def change_own_password():
    data = request_data()
    current_password = password_from(data, "senha_atual")
    new_password = password_from(data, "nova_senha")
    if hmac.compare_digest(current_password, new_password):
        raise ApiError("A nova senha deve ser diferente da senha atual.")

    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT senha FROM usuarios WHERE id = %s", (g.current_user["id"],))
        stored = cursor.fetchone()
        valid_password, _ = verify_password(stored["senha"], current_password)
        if not valid_password:
            raise ApiError("A senha atual está incorreta.", 401)
        cursor.execute(
            "UPDATE usuarios SET senha = %s WHERE id = %s",
            (generate_password_hash(new_password), g.current_user["id"]),
        )
        conn.commit()
    return jsonify({"mensagem": "Senha alterada com sucesso."})


@app.get("/api/usuarios")
@manager_required
def list_users():
    include_inactive = request.args.get("incluir_inativos", "false").lower() == "true"
    query = "SELECT id, nome, tipo_usuario, ativo FROM usuarios"
    if not include_inactive:
        query += " WHERE ativo = TRUE"
    query += " ORDER BY nome"
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(query)
        users = [public_user(user) for user in cursor.fetchall()]
    return jsonify({"usuarios": users})


@app.get("/api/usuarios/desativados")
@manager_required
def list_inactive_users():
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT id, nome, tipo_usuario, ativo FROM usuarios WHERE ativo = FALSE ORDER BY nome"
        )
        users = [public_user(user) for user in cursor.fetchall()]
    return jsonify({"usuarios": users})


@app.post("/api/usuarios")
@manager_required
def create_user():
    data = request_data()
    name = required_text(data, "nome", "Nome")
    password = password_from(data)
    user_type = user_type_from(data)
    try:
        with connection() as conn:
            cursor = conn.cursor(dictionary=True)
            cursor.execute(
                "INSERT INTO usuarios (nome, senha, tipo_usuario, ativo) VALUES (%s, %s, %s, TRUE)",
                (name, generate_password_hash(password), user_type),
            )
            user_id = cursor.lastrowid
            conn.commit()
            cursor.execute(
                "SELECT id, nome, tipo_usuario, ativo FROM usuarios WHERE id = %s", (user_id,)
            )
            user = cursor.fetchone()
    except mysql.connector.IntegrityError:
        raise ApiError("Já existe um usuário com este nome.", 409) from None
    return jsonify({"mensagem": "Usuário cadastrado com sucesso.", "usuario": public_user(user)}), 201


@app.patch("/api/usuarios/<int:user_id>")
@manager_required
def update_user(user_id):
    data = request_data()
    name = optional_text(data, "nome", "Nome")
    password = password_from(data, required=False)
    user_type = user_type_from(data, required=False)
    active = data.get("ativo") if "ativo" in data else None
    if active is not None and not isinstance(active, bool):
        raise ApiError("O campo ativo deve ser verdadeiro ou falso.")
    if all(value is None for value in (name, password, user_type, active)):
        raise ApiError("Informe ao menos um campo para atualização.")
    if user_id == g.current_user["id"] and (user_type == "FUNCIONARIO" or active is False):
        raise ApiError("Não é permitido remover seu próprio acesso de gestor.", 409)
    should_remove_from_queue = active is False or user_type == "GESTOR"

    assignments, values = [], []
    if name is not None:
        assignments.append("nome = %s")
        values.append(name)
    if password is not None:
        assignments.append("senha = %s")
        values.append(generate_password_hash(password))
    if user_type is not None:
        assignments.append("tipo_usuario = %s")
        values.append(user_type)
    if active is not None:
        assignments.append("ativo = %s")
        values.append(active)
    values.append(user_id)

    try:
        with connection() as conn:
            cursor = conn.cursor()
            lock = queue_mutation_lock(conn) if should_remove_from_queue else nullcontext()
            with lock:
                cursor.execute(f"UPDATE usuarios SET {', '.join(assignments)} WHERE id = %s", values)
                if cursor.rowcount == 0:
                    cursor.execute("SELECT 1 FROM usuarios WHERE id = %s", (user_id,))
                    if not cursor.fetchone():
                        raise ApiError("Usuário não encontrado.", 404)
                if should_remove_from_queue:
                    remove_employee_from_queue(cursor, user_id)
                conn.commit()
                read_cursor = conn.cursor(dictionary=True)
                read_cursor.execute(
                    "SELECT id, nome, tipo_usuario, ativo FROM usuarios WHERE id = %s", (user_id,)
                )
                user = read_cursor.fetchone()
    except mysql.connector.IntegrityError:
        raise ApiError("Já existe um usuário com este nome.", 409) from None
    return jsonify({"mensagem": "Usuário atualizado com sucesso.", "usuario": public_user(user)})


@app.post("/api/usuarios/<int:user_id>/reativar")
@manager_required
def reactivate_user(user_id):
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("UPDATE usuarios SET ativo = TRUE WHERE id = %s AND ativo = FALSE", (user_id,))
        if cursor.rowcount == 0:
            cursor.execute("SELECT ativo FROM usuarios WHERE id = %s", (user_id,))
            existing_user = cursor.fetchone()
            if not existing_user:
                raise ApiError("Usuário não encontrado.", 404)
            raise ApiError("Este usuário já está ativo.", 409)
        conn.commit()
        cursor.execute(
            "SELECT id, nome, tipo_usuario, ativo FROM usuarios WHERE id = %s", (user_id,)
        )
        user = cursor.fetchone()
    return jsonify({"mensagem": "Usuário reativado com sucesso.", "usuario": public_user(user)})


@app.delete("/api/usuarios/<int:user_id>/permanente")
@manager_required
def permanently_delete_user(user_id):
    """Exclui fisicamente apenas perfis que já estão desativados."""
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            cursor.execute("SELECT id, ativo FROM usuarios WHERE id = %s FOR UPDATE", (user_id,))
            user = cursor.fetchone()
            if not user:
                raise ApiError("Usuário não encontrado.", 404)
            if user[1]:
                raise ApiError("Apenas usuários desativados podem ser excluídos definitivamente.", 409)

            # Mantém a fila consistente mesmo para registros legados que possam
            # ter permanecido nela após serem desativados.
            remove_employee_from_queue(cursor, user_id)
            cursor.execute("DELETE FROM usuarios WHERE id = %s", (user_id,))
            conn.commit()
    return jsonify({"mensagem": "Usuário excluído definitivamente com sucesso.", "usuario_id": user_id})


@app.delete("/api/usuarios/<int:user_id>")
@manager_required
def deactivate_user(user_id):
    if user_id == g.current_user["id"]:
        raise ApiError("Não é permitido inativar o próprio usuário.", 409)
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            cursor.execute("UPDATE usuarios SET ativo = FALSE WHERE id = %s AND ativo = TRUE", (user_id,))
            if cursor.rowcount == 0:
                raise ApiError("Usuário não encontrado ou já inativo.", 404)
            remove_employee_from_queue(cursor, user_id)
            conn.commit()
    return jsonify({"mensagem": "Usuário inativado com sucesso."})


@app.get("/api/fila")
@authenticated
def get_queue():
    with connection() as conn:
        return jsonify({"fila": queue_rows(conn.cursor())})


@app.get("/api/fila/minha-posicao")
@authenticated
def my_position():
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT posicao, status FROM fila_auditoria WHERE usuario_id = %s",
            (g.current_user["id"],),
        )
        item = cursor.fetchone()
    if not item:
        return jsonify({"na_fila": False, "posicao": None, "mensagem": "Você não está na fila."})
    item["na_fila"] = True
    item["mensagem"] = f'{item["posicao"]}º da fila'
    return jsonify(item)


@app.post("/api/fila/concluir")
@authenticated
def complete_queue():
    with connection() as conn:
        cursor = conn.cursor()
        conn.start_transaction()
        with queue_mutation_lock(conn):
            cursor.execute(
                "SELECT id, usuario_id FROM fila_auditoria WHERE status = 'EM_ANDAMENTO' "
                "ORDER BY posicao LIMIT 1 FOR UPDATE"
            )
            current = cursor.fetchone()
            if not current:
                conn.rollback()
                raise ApiError("Não há auditoria em andamento.", 409)
            current_item_id, current_user_id = current
            if g.current_user["tipo_usuario"] != "GESTOR" and current_user_id != g.current_user["id"]:
                conn.rollback()
                raise ApiError("Apenas o funcionário da vez pode concluir a auditoria.", 403)
            cursor.execute("INSERT INTO auditorias_concluidas (usuario_id) VALUES (%s)", (current_user_id,))
            cursor.execute(
                "SELECT id FROM fila_auditoria WHERE id <> %s ORDER BY posicao, id FOR UPDATE",
                (current_item_id,),
            )
            next_item_ids = [row[0] for row in cursor.fetchall()]
            # O item concluído é anexado ao fim; o primeiro da lista passa a ser a vez atual.
            set_queue_order(cursor, [*next_item_ids, current_item_id])
            conn.commit()
            return jsonify({"mensagem": "Auditoria concluída e funcionário reenfileirado.", "fila": queue_rows(cursor)})


@app.post("/api/fila/adicionar")
@manager_required
def add_to_queue():
    data = request_data()
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            user = find_user_for_queue(cursor, data)
            if not user:
                raise ApiError("Funcionário não encontrado.", 404)
            if user[2] != "FUNCIONARIO":
                raise ApiError("Apenas funcionários podem ser adicionados à fila.", 409)
            cursor.execute("SELECT 1 FROM fila_auditoria WHERE usuario_id = %s", (user[0],))
            if cursor.fetchone():
                raise ApiError("Este funcionário já está na fila.", 409)
            cursor.execute("SELECT COALESCE(MAX(posicao), 0) + 1 FROM fila_auditoria")
            position = cursor.fetchone()[0]
            cursor.execute(
                "INSERT INTO fila_auditoria (usuario_id, posicao, status) VALUES (%s, %s, %s)",
                (user[0], position, "EM_ANDAMENTO" if position == 1 else "AGUARDANDO"),
            )
            conn.commit()
            return jsonify({"fila": queue_rows(cursor)}), 201


@app.post("/api/fila/reordenar")
@manager_required
def reorder_queue():
    data = request_data()
    item_id, action = data.get("id"), str(data.get("acao", "")).upper()
    if not isinstance(item_id, int) or action not in {"SUBIR", "DESCER"}:
        raise ApiError("Informe id numérico e a ação SUBIR ou DESCER.")
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            cursor.execute("SELECT posicao FROM fila_auditoria WHERE id = %s", (item_id,))
            item = cursor.fetchone()
            if not item:
                raise ApiError("Item da fila não encontrado.", 404)
            target = item[0] - 1 if action == "SUBIR" else item[0] + 1
            cursor.execute("SELECT id FROM fila_auditoria WHERE posicao = %s", (target,))
            neighbor = cursor.fetchone()
            if neighbor:
                cursor.execute(
                    "UPDATE fila_auditoria SET posicao = -posicao WHERE id IN (%s, %s)",
                    (item_id, neighbor[0]),
                )
                cursor.execute("UPDATE fila_auditoria SET posicao = %s WHERE id = %s", (target, item_id))
                cursor.execute("UPDATE fila_auditoria SET posicao = %s WHERE id = %s", (item[0], neighbor[0]))
                normalize_queue(cursor)
                conn.commit()
            return jsonify({"fila": queue_rows(cursor)})


@app.post("/api/fila/remover")
@manager_required
def remove_from_queue():
    item_id = request_data().get("id")
    if not isinstance(item_id, int):
        raise ApiError("id numérico é obrigatório.")
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            cursor.execute("DELETE FROM fila_auditoria WHERE id = %s", (item_id,))
            if cursor.rowcount == 0:
                raise ApiError("Item da fila não encontrado.", 404)
            normalize_queue(cursor)
            conn.commit()
            return jsonify({"fila": queue_rows(cursor)})


@app.post("/api/fila/limpar")
@manager_required
def clear_queue():
    with connection() as conn:
        cursor = conn.cursor()
        with queue_mutation_lock(conn):
            cursor.execute("DELETE FROM fila_auditoria")
            conn.commit()
    return jsonify({"fila": []})


@app.get("/api/gestor/ranking")
@authenticated
def ranking():
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        return jsonify({"ranking": metrics_payload(cursor, "dia")["ranking"]})


@app.get("/api/ranking-auditorias")
@authenticated
def audit_ranking():
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        return jsonify({"ranking": audit_ranking_payload(cursor)})


@app.get("/api/gestor/metrics")
@authenticated
def manager_metrics():
    with connection() as conn:
        cursor = conn.cursor(dictionary=True)
        return jsonify(metrics_payload(cursor, request.args.get("periodo", "dia")))


@app.get("/api/gestor/exportar-excel")
@manager_required
def export_excel():
    period, _, _ = metric_period(request.args.get("periodo", "dia"))
    report = build_excel_report(period)
    filename = f"relatorio-auditorias-{period}-{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(
        report,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/gestor/metricas/resetar")
@manager_required
def reset_audit_metrics():
    """Apaga o histórico de auditorias e reinicia o AUTO_INCREMENT no MySQL."""
    data = request_data()
    if data.get("confirmar") is not True:
        raise ApiError("Confirme a limpeza das métricas para continuar.", 400)

    with connection() as conn:
        cursor = conn.cursor()
        # A mesma trava usada na conclusão evita que uma auditoria seja gravada
        # entre a contagem e o reset do histórico.
        with queue_mutation_lock(conn):
            cursor.execute("SELECT COUNT(*) FROM auditorias_concluidas")
            deleted_count = cursor.fetchone()[0]
            # TRUNCATE é o comando MySQL que também reinicia o AUTO_INCREMENT.
            cursor.execute("TRUNCATE TABLE auditorias_concluidas")

    return jsonify(
        {
            "mensagem": "Métricas e ranking foram zerados com sucesso.",
            "auditorias_removidas": deleted_count,
        }
    )


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    app.run(
        host=os.getenv("FLASK_HOST", "127.0.0.1"),
        port=int(os.getenv("PORT", "5000")),
        debug=os.getenv("FLASK_DEBUG", "false").lower() == "true",
    )
